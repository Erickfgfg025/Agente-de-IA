from io import BytesIO
import zipfile

import fitz
import pytest
from openpyxl import load_workbook
from pypdf import PdfWriter

from bot.conversions import available_outputs, convert_file


def blank_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def text_pdf() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 80), "Relatorio de vendas")
    page.insert_text((50, 110), "Janeiro 100")
    result = doc.tobytes()
    doc.close()
    return result


def test_rejects_fake_pdf_named_txt_conversion():
    with pytest.raises(ValueError, match="não parece ser um PDF válido"):
        convert_file(b"isto não é um PDF", "arquivo.pdf", "txt")


def test_converts_txt_to_docx():
    result, name, mime = convert_file("Olá, mundo!".encode(), "saudacao.txt", "docx")
    assert name == "saudacao.docx"
    assert mime.startswith("application/vnd.openxmlformats")
    assert result.startswith(b"PK")


def test_rejects_empty_file():
    with pytest.raises(ValueError, match="arquivo está vazio"):
        convert_file(b"", "vazio.txt", "docx")


def test_pdf_offers_excel_powerpoint_and_jpg():
    assert {"xlsx", "pptx", "jpg"}.issubset(available_outputs("relatorio.pdf"))


def test_pdf_to_excel_creates_valid_xlsx():
    result, filename, mime = convert_file(text_pdf(), "relatorio.pdf", "xlsx")
    assert filename == "relatorio.xlsx"
    assert mime.endswith("spreadsheetml.sheet")
    workbook = load_workbook(BytesIO(result), read_only=True)
    assert workbook.sheetnames == ["Página 1"]
    assert workbook["Página 1"]["A1"].value == "Relatorio de vendas"
    workbook.close()


def test_xlsx_validation_rejects_arbitrary_zip():
    invalid = BytesIO()
    with zipfile.ZipFile(invalid, "w") as archive:
        archive.writestr("random.txt", "not an office file")
    with pytest.raises(ValueError, match="documento .xlsx válido"):
        convert_file(invalid.getvalue(), "bad.xlsx", "pdf")
