from io import BytesIO
from pathlib import Path
import shutil
import subprocess
import tempfile
import zipfile

import fitz
import pdfplumber
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader

from bot.documents import make_docx, make_pdf
from bot.pdf_tools import pdf_to_jpg


SUPPORTED_INPUTS = {
    ".pdf", ".docx", ".txt", ".png", ".jpg", ".jpeg", ".webp", ".xlsx", ".pptx"
}

MIME_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "jpg": "image/jpeg",
}


def available_outputs(filename: str) -> list[str]:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return ["txt", "docx", "xlsx", "pptx", "jpg"]
    if ext == ".docx":
        return ["txt"]
    if ext == ".txt":
        return ["pdf", "docx"]
    if ext in {".png", ".jpg", ".jpeg", ".webp"}:
        return ["pdf"]
    if ext in {".xlsx", ".pptx"}:
        return ["pdf"]
    return []


def _validate_file_signature(data: bytes, ext: str) -> None:
    if not data:
        raise ValueError("O arquivo está vazio.")
    if ext == ".pdf":
        if b"%PDF-" not in data[:1024]:
            raise ValueError("O conteúdo enviado não parece ser um PDF válido.")
    elif ext in {".docx", ".xlsx", ".pptx"}:
        try:
            with zipfile.ZipFile(BytesIO(data)) as archive:
                members = set(archive.namelist())
                expected = {
                    ".docx": {"[Content_Types].xml", "word/document.xml"},
                    ".xlsx": {"[Content_Types].xml", "xl/workbook.xml"},
                    ".pptx": {"[Content_Types].xml", "ppt/presentation.xml"},
                }[ext]
                if not expected.issubset(members):
                    raise ValueError(f"O arquivo não parece ser um documento {ext} válido.")
                bad_member = archive.testzip()
                if bad_member:
                    raise ValueError("O documento Office está danificado.")
        except zipfile.BadZipFile as exc:
            raise ValueError(f"O conteúdo enviado não parece ser um documento {ext} válido.") from exc
    elif ext in {".png", ".jpg", ".jpeg", ".webp"}:
        try:
            with Image.open(BytesIO(data)) as image:
                image.verify()
        except Exception as exc:
            raise ValueError("O conteúdo enviado não parece ser uma imagem válida.") from exc
    elif ext == ".txt":
        try:
            data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("O TXT precisa estar codificado em UTF-8.") from exc


def _safe_cell(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _pdf_to_xlsx(data: bytes, filename: str) -> tuple[bytes, str, str]:
    reader = PdfReader(BytesIO(data))
    if reader.is_encrypted:
        raise ValueError("Este PDF está protegido por senha.")
    if len(reader.pages) > 100:
        raise ValueError("A conversão de PDF para Excel aceita até 100 páginas.")
    workbook = Workbook()
    workbook.remove(workbook.active)
    extracted_any_table = False
    try:
        with pdfplumber.open(BytesIO(data)) as pdf:
            for page_index, page in enumerate(pdf.pages, 1):
                worksheet = workbook.create_sheet(f"Página {page_index}")
                tables = page.extract_tables() or []
                if tables:
                    extracted_any_table = True
                    for table_index, table in enumerate(tables):
                        if table_index:
                            worksheet.append([])
                            worksheet.append([f"Tabela {table_index + 1}"])
                        for row in table:
                            worksheet.append([_safe_cell(cell or "") for cell in (row or [])])
                else:
                    text = page.extract_text() or ""
                    for line in text.splitlines():
                        worksheet.append([_safe_cell(line)])
        if not extracted_any_table and all(
            ws.max_row == 1 and ws["A1"].value is None for ws in workbook.worksheets
        ):
            raise ValueError("Não encontrei texto ou tabelas extraíveis. Se o PDF for digitalizado, execute OCR primeiro.")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Não consegui extrair tabelas deste PDF. Se ele for digitalizado, execute OCR primeiro.") from exc

    output = BytesIO()
    workbook.save(output)
    stem = Path(filename).stem or "documento"
    return output.getvalue(), f"{stem}.xlsx", MIME_TYPES["xlsx"]


def _pdf_to_pptx(data: bytes, filename: str) -> tuple[bytes, str, str]:
    stem = Path(filename).stem or "documento"
    presentation = Presentation()
    slide_width = presentation.slide_width
    slide_height = presentation.slide_height
    blank_layout = presentation.slide_layouts[6]
    try:
        with fitz.open(stream=data, filetype="pdf") as pdf:
            if len(pdf) > 100:
                raise ValueError("A conversão de PDF para PowerPoint aceita até 100 páginas.")
            if len(pdf) == 0:
                raise ValueError("O PDF não contém páginas.")
            for page in pdf:
                pixmap = page.get_pixmap(matrix=fitz.Matrix(1.3, 1.3), alpha=False)
                if pixmap.width * pixmap.height > 40_000_000:
                    raise ValueError("Uma página tem resolução excessiva para esta conversão.")
                stream = BytesIO(pixmap.tobytes("png"))
                slide = presentation.slides.add_slide(blank_layout)
                page_ratio = pixmap.width / pixmap.height
                canvas_ratio = slide_width / slide_height
                if page_ratio > canvas_ratio:
                    width = slide_width
                    height = int(slide_width / page_ratio)
                else:
                    height = slide_height
                    width = int(slide_height * page_ratio)
                left = int((slide_width - width) / 2)
                top = int((slide_height - height) / 2)
                slide.shapes.add_picture(stream, left, top, width=width, height=height)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Não consegui criar a apresentação a partir deste PDF.") from exc

    output = BytesIO()
    presentation.save(output)
    return output.getvalue(), f"{stem}.pptx", MIME_TYPES["pptx"]


def _office_to_pdf(data: bytes, filename: str) -> tuple[bytes, str, str]:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".xlsx", ".pptx"}:
        raise ValueError("Este conversor aceita Excel (.xlsx) ou PowerPoint (.pptx).")
    stem = Path(filename).stem or "documento"
    if shutil.which("soffice") is None:
        raise ValueError("O conversor Office não está instalado no servidor. Verifique o deploy Docker.")
    try:
        with tempfile.TemporaryDirectory(prefix="documento-pronto-") as temp_dir:
            root = Path(temp_dir)
            source = root / f"{stem}{suffix}"
            output_dir = root / "output"
            profile_dir = root / "lo-profile"
            output_dir.mkdir()
            source.write_bytes(data)
            result = subprocess.run(
                [
                    "soffice",
                    f"-env:UserInstallation={profile_dir.as_uri()}",
                    "--headless",
                    "--convert-to", "pdf",
                    "--outdir", str(output_dir),
                    str(source),
                ],
                check=False,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                text=True,
                timeout=90,
            )
            target = output_dir / f"{stem}.pdf"
            if result.returncode != 0 or not target.exists():
                raise ValueError("O LibreOffice não conseguiu converter este arquivo. Verifique se ele está íntegro.")
            pdf_bytes = target.read_bytes()
            if not pdf_bytes.startswith(b"%PDF-"):
                raise ValueError("A conversão não gerou um PDF válido.")
            return pdf_bytes, f"{stem}.pdf", MIME_TYPES["pdf"]
    except subprocess.TimeoutExpired as exc:
        raise ValueError("A conversão demorou demais. Tente um arquivo menor.") from exc
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Não consegui converter o arquivo Office para PDF.") from exc


def convert_file(data: bytes, filename: str, output_format: str) -> tuple[bytes, str, str]:
    ext = Path(filename).suffix.lower()
    stem = Path(filename).stem or "arquivo"
    _validate_file_signature(data, ext)

    if ext == ".pdf":
        if output_format == "xlsx":
            return _pdf_to_xlsx(data, filename)
        if output_format == "pptx":
            return _pdf_to_pptx(data, filename)
        if output_format == "jpg":
            return pdf_to_jpg(data, filename)
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            raise ValueError("Este PDF está protegido por senha.")
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages).strip()
        if not text:
            raise ValueError("Este PDF não contém texto extraível. Se for digitalizado, use Ferramentas de PDF > OCR.")
        if output_format == "txt":
            return text.encode("utf-8"), f"{stem}.txt", MIME_TYPES["txt"]
        if output_format == "docx":
            return make_docx(text, stem), f"{stem}.docx", MIME_TYPES["docx"]

    elif ext == ".docx" and output_format == "txt":
        doc = Document(BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs).strip()
        for table in doc.tables:
            for row in table.rows:
                text += "\n" + "\t".join(cell.text for cell in row.cells)
        return text.encode("utf-8"), f"{stem}.txt", MIME_TYPES["txt"]

    elif ext == ".txt":
        text = data.decode("utf-8-sig")
        if output_format == "pdf":
            return make_pdf(text, stem), f"{stem}.pdf", MIME_TYPES["pdf"]
        if output_format == "docx":
            return make_docx(text, stem), f"{stem}.docx", MIME_TYPES["docx"]

    elif ext in {".png", ".jpg", ".jpeg", ".webp"} and output_format == "pdf":
        with Image.open(BytesIO(data)) as image:
            if image.width * image.height > 40_000_000:
                raise ValueError("A imagem tem resolução muito alta para converter com segurança.")
            if image.mode not in ("RGB", "L"):
                image = image.convert("RGB")
            output = BytesIO()
            image.save(output, format="PDF", resolution=150.0)
            return output.getvalue(), f"{stem}.pdf", MIME_TYPES["pdf"]

    elif ext in {".xlsx", ".pptx"} and output_format == "pdf":
        return _office_to_pdf(data, filename)

    raise ValueError("Esta combinação de formato não é compatível.")
