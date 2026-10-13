from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.shared import Inches
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.enums import TA_LEFT
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase import pdfmetrics
from reportlab.lib import colors
from xml.sax.saxutils import escape


def safe_filename(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9À-ÿ _-]+", "", value).strip()
    value = re.sub(r"\s+", "_", value)[:60]
    return value or "documento"


def make_docx(text: str, title: str = "Documento") -> bytes:
    doc = Document()
    doc.core_properties.title = title
    doc.add_heading(title, level=1)
    for paragraph in text.splitlines():
        doc.add_paragraph(paragraph)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def make_pdf(text: str, title: str = "Documento") -> bytes:
    output = BytesIO()
    styles = getSampleStyleSheet()
    title_style = styles["Title"]
    body_style = styles["BodyText"]
    body_style.leading = 14
    story = [Paragraph(escape(title), title_style), Spacer(1, 0.5 * cm)]
    for line in text.splitlines():
        story.append(Paragraph(escape(line) if line else "&#160;", body_style))
        story.append(Spacer(1, 0.12 * cm))
    pdf = SimpleDocTemplate(
        output, pagesize=A4, rightMargin=2 * cm, leftMargin=2 * cm,
        topMargin=2 * cm, bottomMargin=2 * cm, title=title,
    )
    pdf.build(story)
    return output.getvalue()


def extract_text_from_pdf(data: bytes) -> str:
    reader = PdfReader(BytesIO(data))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages).strip()
